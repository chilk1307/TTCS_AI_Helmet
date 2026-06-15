import os
import threading
from typing import List, Optional
from datetime import datetime

import openpyxl
from openpyxl import Workbook

from app.core import config
from app.core.logger import get_logger
from app.services.inference.types import ViolationRecord

logger = get_logger(__name__)

HEADERS = ["id", "ten", "bien_so", "minh_chung", "thoi_gian", "frame_index", "confidence", "video_id"]


class SheetWriter:
    """Write violations to Excel file. Thread-safe."""

    def __init__(self):
        self._lock = threading.Lock()
        os.makedirs(config.VIOLATIONS_DIR, exist_ok=True)
        self._init_file()

    def _init_file(self):
        if not os.path.exists(config.VIOLATIONS_FILE):
            wb = Workbook()
            ws = wb.active
            ws.title = "Violations"
            ws.append(HEADERS)
            wb.save(config.VIOLATIONS_FILE)
            logger.info(f"Created violations file: {config.VIOLATIONS_FILE}")

    def save_violation(self, record: ViolationRecord):
        with self._lock:
            try:
                wb = openpyxl.load_workbook(config.VIOLATIONS_FILE)
                ws = wb.active
                ws.append([
                    record.id,
                    record.ten,
                    record.bien_so,
                    record.minh_chung,
                    record.thoi_gian,
                    record.frame_index,
                    record.confidence,
                    record.video_id,
                ])
                wb.save(config.VIOLATIONS_FILE)
                logger.info(f"Violation saved: {record.id} | plate={record.bien_so}")
            except Exception as e:
                logger.error(f"Error saving violation: {e}")

    def list_violations(self, video_id: Optional[str] = None) -> List[dict]:
        with self._lock:
            try:
                if not os.path.exists(config.VIOLATIONS_FILE):
                    return []
                wb = openpyxl.load_workbook(config.VIOLATIONS_FILE, read_only=True)
                ws = wb.active
                rows = list(ws.iter_rows(min_row=2, values_only=True))
                violations = []
                for row in rows:
                    if len(row) < len(HEADERS):
                        continue
                    record = dict(zip(HEADERS, row))
                    if video_id and record.get("video_id") != video_id:
                        continue
                    violations.append(record)
                return violations
            except Exception as e:
                logger.error(f"Error reading violations: {e}")
                return []
